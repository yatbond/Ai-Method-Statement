"""
Update PRD, UX Flow, and Prompt Library docs to reflect Clerk auth migration.
Run from: /home/yatbond/projects/Ai-Method-Statement
"""
import copy
import docx
from docx import Document
from docx.oxml.ns import qn


def replace_in_para(para, old, new):
    """Replace text in a paragraph across all runs, preserving style of first run."""
    full = "".join(r.text for r in para.runs)
    if old not in full:
        return False
    updated = full.replace(old, new)
    if para.runs:
        para.runs[0].text = updated
        for r in para.runs[1:]:
            r.text = ""
    return True


def replace_in_cell(cell, old, new):
    """Replace text anywhere in a table cell."""
    changed = False
    for para in cell.paragraphs:
        if replace_in_para(para, old, new):
            changed = True
    return changed


def replace_in_table(table, old, new):
    for row in table.rows:
        for cell in row.cells:
            replace_in_cell(cell, old, new)


def replace_everywhere(doc, old, new):
    """Replace in all paragraphs and all table cells."""
    for para in doc.paragraphs:
        replace_in_para(para, old, new)
    for table in doc.tables:
        replace_in_table(table, old, new)


# ─────────────────────────────────────────────
# PRD
# ─────────────────────────────────────────────
print("Updating PRD...")
prd = Document("Documentations/AI_Method_Statement_Studio_PRD_v3.docx")

replacements_prd = [
    # About this revision
    (
        "SSO replaces hosted-only auth assumption.",
        "Clerk (email + password) replaces SAML/SSO for the initial internal deployment.",
    ),
    # Reference implementation (tech stack)
    (
        "SAML SSO into the company IdP;",
        "Clerk authentication (email + password);",
    ),
    # Admin role description
    (
        "SSO authentication and project/role-based access.",
        "Clerk authentication and project/role-based access.",
    ),
    # Phase 0 open decisions checklist
    (
        "SSO provider and IdP protocol confirmed (Open Decision §13).",
        "Clerk authentication configured; API keys obtained and webhook endpoint registered.",
    ),
    # Phase 0 done criteria
    (
        "Engineer can log in via SSO, create a project, select a trade, and reach an empty method statement shell.",
        "Engineer can log in via Clerk (email + password), create a project, select a trade, and reach an empty method statement shell.",
    ),
    # User journey table step 1
    (
        "Authenticate via SSO; show project list scoped to user's access.",
        "Authenticate via Clerk (email + password); show project list scoped to user's access.",
    ),
    # NFR REQ-NFR-SEC-001
    (
        "Authentication MUST use the company's existing SSO (Azure AD or equivalent), not a hosted email/password service.",
        (
            "Authentication uses Clerk (email + password). Users authenticate directly within the app; "
            "session management is handled by Clerk middleware. Clerk webhooks sync user identity to the "
            "application database on user.created and user.updated events. "
            "SSO integration with the company identity provider (Azure AD) MAY be added in a later release "
            "if required by IT Security policy."
        ),
    ),
    # Tech stack table — Authentication row
    (
        "SSO integration (SAML / OIDC) with the company identity provider. Hosted-only auth services are not acceptable.",
        (
            "Clerk authentication (email + password). Sign-in, session management, and sign-out handled by "
            "the Clerk SDK (@clerk/nextjs). User identity synced to PostgreSQL via Clerk webhook "
            "(user.created / user.updated → database upsert)."
        ),
    ),
    # Phase 0 deliverables table
    (
        "SSO authentication integrated with company IdP (SAML/OIDC). Login, session, logout.",
        (
            "Clerk authentication integrated. Login (email + password via Clerk SignIn component), "
            "session management (Clerk clerkMiddleware), sign-out (Clerk UserButton). "
            "User record upserted to PostgreSQL via webhook on user.created / user.updated."
        ),
    ),
    # Open decisions table
    (
        "Identity provider and SSO protocol",
        "Authentication provider",
    ),
    (
        "IT Security\t MVP build start",
        "IT Security\t DECIDED — Clerk (email + password) for internal deployment",
    ),
    # Catch remaining generic "IT Security | MVP build start" variant
    (
        "IT Security | MVP build start",
        "IT Security | DECIDED — Clerk (email + password) for internal deployment",
    ),
]

for old, new in replacements_prd:
    replace_everywhere(prd, old, new)

prd.save("Documentations/AI_Method_Statement_Studio_PRD_v3.docx")
print("  PRD saved.")


# ─────────────────────────────────────────────
# UX FLOW
# ─────────────────────────────────────────────
print("Updating UX Flow...")
ux = Document("Documentations/UX_Flow_Method_Statement_Studio.docx")

replacements_ux = [
    # Screen heading
    (
        "S-01 — Login / SSO handoff",
        "S-01 — Login",
    ),
    # Screen inventory table
    (
        "Login / SSO handoff",
        "Login",
    ),
    # S-01 Purpose
    (
        "Authenticate the user via the company SSO provider. No credentials are entered in the app.",
        "Authenticate the user via Clerk. User enters their email and password directly in the app.",
    ),
    # S-01 Layout
    (
        "Single centred card on a neutral background. Company logo top-centre. One primary button. No form fields.",
        (
            "Single centred card on a neutral background. Company logo and app title top-centre. "
            "Clerk SignIn component renders below (email field, password field, Sign in button, "
            "forgot-password link). No external redirect."
        ),
    ),
    # S-01 Elements table — button row
    (
        "Button (primary): \"Sign in with [Company] account\"",
        (
            "Clerk SignIn component: email address field (type=email, required), "
            "password field (type=password, required), Sign in button (primary). "
            "Forgot password link beneath the button."
        ),
    ),
    # S-01 Actions table — success flow
    (
        "Clicks \"Sign in with [Company] account\"",
        "Submits email + password",
    ),
    (
        "Redirect to company IdP (Azure AD / Okta). On successful SSO callback, create or update user session. Apply role from IdP group claim.",
        (
            "Clerk authenticates the credentials. On success: Clerk session cookie set; "
            "Clerk fires user.created / user.updated webhook → PostgreSQL user record upserted. "
            "App redirects to S-02."
        ),
    ),
    # S-01 Actions table — failure flow
    (
        "SSO fails or times out",
        "Wrong email or password",
    ),
    (
        "Return to S-01 with error message: \"Sign-in failed. Please try again or contact IT support.\"",
        "Clerk renders inline error below the form: \"Incorrect email or password.\" No redirect.",
    ),
    # S-01 edge case — no assigned role
    (
        "Apply role from IdP group claim.",
        "User record created in database with default role (ENGINEER) on first sign-in.",
    ),
    # Header bar description
    (
        "Header bar with user avatar, role badge, and sign-out.",
        "Header bar with user name, Clerk UserButton (avatar + sign-out dropdown).",
    ),
    # S-23 Admin — user management purpose
    (
        "Manage users, roles, and project access. Sync with company IdP; override roles where needed.",
        "Manage users, roles, and project access. Users are provisioned via Clerk; roles are managed within the application.",
    ),
    # S-23 User list columns — IdP group → Clerk User ID
    (
        "[Name] | [Email] | [System role] | [IdP group] | [Last login] | [Status: Active/Inactive]",
        "[Name] | [Email] | [System role] | [Clerk User ID] | [Last login] | [Status: Active/Inactive]",
    ),
    # S-23 role sync description
    (
        "└ Assigned system role (synced from IdP; overridable by Admin)",
        "└ Assigned system role (set in application; overridable by Admin)",
    ),
    # Version line
    (
        "Version 1.0  |  Companion to PRD v2.0",
        "Version 1.1  |  Companion to PRD v3.0",
    ),
]

for old, new in replacements_ux:
    replace_everywhere(ux, old, new)

ux.save("Documentations/UX_Flow_Method_Statement_Studio.docx")
print("  UX Flow saved.")


# ─────────────────────────────────────────────
# PROMPT LIBRARY — version bump only
# ─────────────────────────────────────────────
print("Updating Prompt Library...")
pl = Document("Documentations/Prompt_Library_Method_Statement_Studio.docx")

replacements_pl = [
    (
        "Version 1.0  |  Companion to PRD v2.0  |  Requires AI safety reviewer sign-off before production use",
        "Version 1.0  |  Companion to PRD v3.0  |  Requires AI safety reviewer sign-off before production use",
    ),
]

for old, new in replacements_pl:
    replace_everywhere(pl, old, new)

pl.save("Documentations/Prompt_Library_Method_Statement_Studio.docx")
print("  Prompt Library saved.")

print("\nAll done.")
